"""Generate Outlook-friendly full and short academic email signatures.

Keep academic role facts synchronized with the canonical public profile.
The signatures start with a single '---' line in normal text style and retain the exact motto and upright emoji.
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
FULL_OUT = HERE / "Mineo_Raffaele_EmailSignature.docx"
SHORT_OUT = HERE / "Mineo_Raffaele_EmailSignature_Short.docx"
FONT = "Segoe UI"
NAME_COLOR = "404040"
TEXT_COLOR = "606060"
LINK_COLOR = "365F7D"
EMAIL = "raffaele.mineo@phd.unict.it"
WEBSITE = "https://ingraffaelemineo.github.io/"
MOTTO = "My Life is The Research, and I am truly lucky to be able to live it to the fullest"

# Conference shorthand keeps the signature compact; complete role names remain
# recorded on the website, in the academic CV, and in AcademicActivities.
FULL_IDENTITY = [
    [("text", "---")],  # Required 9 pt #606060 first line, followed by Shift+Enter.
    [("text", "Dr. "), ("name", "Raffaele Mineo"), ("text", ", Eng. Ph.D.")],
    [("text", "Italian National PhD in Artificial Intelligence - Health and Life Sciences | Member, IEEE")],
    [("text", "Research Fellow - PeRCeiVe.AI Lab | University of Catania")],
    [("text", "Department of Industrial, Electrical, Electronic and Computer Engineering (DIEEI)")],
]

# Fixed contiguous category order: Organizing Committee, Associate Editor,
# Guest Editor (Lead before regular), Area Chair. Within each category,
# use descending standing per the repository's GGS/ICORE and SCImago rules.
FULL_ROLES = [
    [("text", "Organizing Committee - IEEE ISBI 2027, Associate to Technical Program Chairs")],
    [("text", "Associate Editor - Frontiers in Communication")],
    [("text", "Lead Guest Editor - Springer Cognitive Computation")],
    [("text", "Area Chair - MICCAI - IEEE ICASSP - IJCNN - IEEE FG")],
]

FULL_CONTACT = [
    [("text", "Polo Tecnologico, Via Santa Sofia 102, 95123 Catania, Italy")],
    [("text", "Email: "), ("link", EMAIL, f"mailto:{EMAIL}"),
     ("text", " | Web: "), ("link", "ingraffaelemineo.github.io", WEBSITE)],
    [("motto", MOTTO), ("emoji", " 🙃")],
]

SHORT_LINES = [
    FULL_IDENTITY[0],  # The identical --- prefix for replies and forwards.
    FULL_IDENTITY[1],
    [("text", "Research Fellow | PeRCeiVe.AI Lab | University of Catania")],
    [("link", EMAIL, f"mailto:{EMAIL}"),
     ("text", " | "), ("link", "ingraffaelemineo.github.io", WEBSITE)],
    [("motto", MOTTO), ("emoji", " 🙃")],
]


def make_document(title):
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.32)
    section.bottom_margin = Inches(0.32)
    section.left_margin = Inches(0.44)
    section.right_margin = Inches(0.44)

    normal = document.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9)
    normal.font.bold = False
    normal.font.italic = False
    normal.font.underline = False
    normal.font.color.rgb = RGBColor.from_string(TEXT_COLOR)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    document.core_properties.title = title
    document.core_properties.subject = "Outlook-compatible academic email signature"
    document.core_properties.author = "Raffaele Mineo"
    return document


def new_paragraph(document, after_pt=0):
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after_pt)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    return paragraph


def add_text(paragraph, value, is_name=False):
    run = paragraph.add_run(value)
    run.font.name = FONT
    run.font.size = Pt(10 if is_name else 9)
    run.font.bold = bool(is_name)
    run.font.italic = False
    run.font.underline = False
    run.font.color.rgb = RGBColor.from_string(NAME_COLOR if is_name else TEXT_COLOR)
    return run


def add_hyperlink(paragraph, label, target):
    relation = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + key), FONT)
    properties.append(fonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), LINK_COLOR)
    properties.append(color)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    properties.append(size)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "none")
    properties.append(underline)
    run.append(properties)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_soft_lines(paragraph, lines):
    for index, fragments in enumerate(lines):
        if index:
            add_text(paragraph, "").add_break()  # Shift+Enter, no new paragraph.
        for part in fragments:
            kind = part[0]
            if kind == "text":
                add_text(paragraph, part[1])
            elif kind == "name":
                add_text(paragraph, part[1], is_name=True)
            elif kind == "link":
                add_hyperlink(paragraph, part[1], part[2])
            elif kind == "motto":
                add_text(paragraph, part[1]).font.italic = True
            elif kind == "emoji":
                # Upright smile, with ordinary text styling and no emphasis.
                add_text(paragraph, part[1]).font.italic = False
            else:
                raise ValueError(f"Unknown signature fragment: {kind}")


def build_full():
    document = make_document("Raffaele Mineo - Full email signature")
    add_soft_lines(new_paragraph(document, after_pt=3.5), FULL_IDENTITY)
    add_soft_lines(new_paragraph(document), FULL_ROLES)
    add_soft_lines(new_paragraph(document), FULL_CONTACT)
    document.save(FULL_OUT)


def build_short():
    document = make_document("Raffaele Mineo - Short email signature")
    add_soft_lines(new_paragraph(document), SHORT_LINES)
    document.save(SHORT_OUT)


if __name__ == "__main__":
    build_full()
    build_short()
    print(f"Generated: {FULL_OUT.name}, {SHORT_OUT.name}")
