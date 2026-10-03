from pathlib import Path
import json
import re
import unicodedata

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "Mineo_Raffaele_PublicationsMetadata.xlsx"
DOCX = ROOT / "Mineo_Raffaele_PublicationsList.docx"
MANIFEST = ROOT / "Mineo_Raffaele_PublicationFiles.md"
ARCHIVE = ROOT / "publications"
PROFILE_DATA = ROOT.parent / "profile-data.js"

FORBIDDEN_TITLE = "Memory-Augmented Prompt Tuning for Self-Supervised Continual Learning in Medical Imaging"
FORBIDDEN_TOKEN = "MAPT"


def clean_xlsx() -> bool:
    wb = load_workbook(XLSX)
    ws = wb["Foglio1"]
    changed = False

    rows_to_delete = []
    for row in range(2, ws.max_row + 1):
        title = str(ws.cell(row=row, column=1).value or "")
        if title == FORBIDDEN_TITLE or FORBIDDEN_TOKEN in title:
            rows_to_delete.append(row)

    for row in reversed(rows_to_delete):
        ws.delete_rows(row, 1)
        changed = True

    last_row = 1
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value not in (None, ""):
            last_row = row
            if ws.row_dimensions[row].height != 15:
                ws.row_dimensions[row].height = 15
                changed = True

    if "Tabella1" in ws.tables:
        expected_ref = f"A1:K{last_row}"
        if ws.tables["Tabella1"].ref != expected_ref:
            ws.tables["Tabella1"].ref = expected_ref
            changed = True

        style = ws.tables["Tabella1"].tableStyleInfo
        if style is not None:
            if style.name != "TableStyleMedium14":
                style.name = "TableStyleMedium14"
                changed = True
            if not style.showRowStripes:
                style.showRowStripes = True
                changed = True

    if changed:
        wb.save(XLSX)
    return changed


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def clean_docx() -> bool:
    doc = Document(DOCX)
    paragraphs = list(doc.paragraphs)
    target_indexes = [
        i
        for i, p in enumerate(paragraphs)
        if FORBIDDEN_TITLE in p.text or FORBIDDEN_TOKEN in p.text
    ]
    if not target_indexes:
        return False

    # PublicationList uses three consecutive paragraphs per record:
    # title/year, authors, and venue/type/link.
    to_remove = set()
    for index in target_indexes:
        to_remove.update(range(index, min(index + 3, len(paragraphs))))

    for index in sorted(to_remove, reverse=True):
        remove_paragraph(paragraphs[index])

    doc.save(DOCX)
    return True



def normalize_title(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("‐", "-").replace("‑", "-").replace("–", "-").replace("—", "-")
    return re.sub(r"\s+", " ", value).strip().casefold()


def verify_repository_consistency() -> None:
    wb = load_workbook(XLSX, read_only=False)
    ws = wb["Foglio1"]
    xlsx_titles = [
        str(ws.cell(row=row, column=1).value)
        for row in range(2, ws.max_row + 1)
        if ws.cell(row=row, column=1).value not in (None, "")
    ]

    assert all(FORBIDDEN_TITLE not in title and FORBIDDEN_TOKEN not in title for title in xlsx_titles), "MAPT remains in XLSX"

    table = ws.tables.get("Tabella1")
    assert table is not None, "Tabella1 is missing"
    expected_ref = f"A1:K{len(xlsx_titles) + 1}"
    assert table.ref == expected_ref, f"Tabella1 ref is {table.ref}, expected {expected_ref}"
    assert table.tableStyleInfo is not None and table.tableStyleInfo.name == "TableStyleMedium14"
    assert bool(table.tableStyleInfo.showRowStripes), "Tabella1 banded rows are disabled"
    for row in range(2, len(xlsx_titles) + 2):
        assert ws.row_dimensions[row].height == 15, f"Row {row} height is not 15 pt"

    doc = Document(DOCX)
    doc_text = "\n".join(p.text for p in doc.paragraphs)
    assert FORBIDDEN_TITLE not in doc_text and FORBIDDEN_TOKEN not in doc_text, "MAPT remains in DOCX"

    raw = PROFILE_DATA.read_text(encoding="utf-8").strip()
    prefix = "window.PROFILE_DATA = "
    assert raw.startswith(prefix), "Unexpected profile-data.js format"
    payload = raw[len(prefix):].rstrip(";\n ")
    profile = json.loads(payload)
    profile_titles = [str(item["title"]) for item in profile["publications"]]
    assert FORBIDDEN_TITLE not in profile_titles, "MAPT remains in profile-data.js"

    archive_files = sorted(p.name for p in ARCHIVE.iterdir() if p.is_file())
    manifest_text = MANIFEST.read_text(encoding="utf-8")
    manifest_files = []
    for line in manifest_text.splitlines():
        parts = line.split(chr(96))
        for part in parts:
            if part.lower().endswith((".pdf", ".zip")):
                manifest_files.append(part)
    manifest_files = sorted(dict.fromkeys(manifest_files))

    assert len(set(map(normalize_title, xlsx_titles))) == len(xlsx_titles), "Duplicate XLSX publication titles"
    assert len(profile_titles) == len(xlsx_titles), f"Profile/XLSX count mismatch: {len(profile_titles)} vs {len(xlsx_titles)}"
    assert len(archive_files) == len(xlsx_titles), f"Archive/XLSX count mismatch: {len(archive_files)} vs {len(xlsx_titles)}"
    assert archive_files == manifest_files, "Publication manifest does not match archived PDF/ZIP set"

    print(f"Canonical publication/output records: {len(xlsx_titles)}")
    print(f"Archived PDF/ZIP files: {len(archive_files)}")
    print(f"Manifested PDF/ZIP files: {len(manifest_files)}")
    print(f"Tabella1 range: {table.ref}")
    print("MAPT absent from XLSX, DOCX and profile-data.js")


def main() -> None:
    xlsx_changed = clean_xlsx()
    docx_changed = clean_docx()
    print(f"XLSX changed: {xlsx_changed}")
    print(f"DOCX changed: {docx_changed}")
    verify_repository_consistency()


if __name__ == "__main__":
    main()
