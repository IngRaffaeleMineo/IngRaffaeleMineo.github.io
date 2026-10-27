"""Regression checks for Outlook-friendly full and short DOCX signatures."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

HERE = Path(__file__).resolve().parent
FULL = HERE / "Mineo_Raffaele_EmailSignature.docx"
SHORT = HERE / "Mineo_Raffaele_EmailSignature_Short.docx"


def all_text(doc):
    return "\n".join("".join(t.text or "" for t in p._p.iter(qn("w:t"))) for p in doc.paragraphs)


def check(path, n_paragraphs, n_soft_breaks, n_links, full):
    doc = Document(path)
    assert len(doc.paragraphs) == n_paragraphs, f"Wrong paragraph count: {path}"
    assert not doc.tables and not doc.inline_shapes, "No tables/images in Outlook signatures"
    assert sum(len(list(p._p.iter(qn("w:br")))) for p in doc.paragraphs) == n_soft_breaks
    assert all("".join(t.text or "" for t in p._p.iter(qn("w:t"))).strip() for p in doc.paragraphs), "Never insert blank paragraphs"
    text = all_text(doc)
    assert "Raffaele Mineo" in text and "Eng. Ph.D." in text
    if full:
        assert "Italian National PhD in Artificial Intelligence - Health and Life Sciences" in text
    else:
        assert "Italian National PhD in Artificial Intelligence - Health and Life Sciences" not in text
    assert "Research Fellow: PeRCeiVe.AI Lab" in text and "University of Catania" in text
    assert text.count("---") == 1, "A single three-hyphen prefix is required"
    assert doc.paragraphs[0].text.startswith("---\nDr. "), "The first line must be exactly --- followed by a soft break before the signature"
    first_run = doc.paragraphs[0].runs[0]
    assert first_run.text == "---", "First DOCX run must consist of three literal hyphens"
    assert first_run.font.name == "Segoe UI" and first_run.font.size == Pt(9)
    assert first_run.font.bold is False and not first_run.font.italic
    assert first_run.font.color.rgb is not None and str(first_run.font.color.rgb) == "606060"
    motto = "My Life is The Research, and I am truly lucky to be able to live it to the fullest"
    assert text.count(motto) == 1 and text.count("🙃") == 1
    assert text.index(motto) > text.index("ingraffaelemineo.github.io")
    assert text.index("🙃") > text.index(motto)
    assert "EEAI" not in text
    if full:
        for phrase in ("Member, IEEE", "Organizing Committee", "Paper Selection Committee: IPMI 2027", "Associate Editor", "Lead Guest Editor", "Area Chair", "MICCAI", "ICASSP", "IJCNN", "FG", "ISBI 2027", "DIEEI", "Polo Tecnologico"):
            assert phrase in text, f"Full signature missing: {phrase}"
        assert text.index("Organizing Committee") < text.index("Paper Selection Committee: IPMI 2027") < text.index("Associate Editor") < text.index("Lead Guest Editor") < text.index("Area Chair")
        assert text.count("Paper Selection Committee: IPMI 2027") == 1, "IPMI committee appointment must appear exactly once"
        assert text.index("MICCAI") < text.index("ICASSP") < text.index("IJCNN") < text.index("FG")
        assert [p.paragraph_format.space_after for p in doc.paragraphs] == [Pt(3.5), Pt(3.5), Pt(3.5), Pt(0)], "The full signature needs three equal 3.5pt inter-block gaps"
        assert doc.paragraphs[1].text.startswith("Organizing Committee: IEEE ISBI 2027")
        assert "Area Chair: MICCAI - IEEE ICASSP - IJCNN - IEEE FG" in text
        for role_type in ("Organizing Committee", "Paper Selection Committee", "Associate Editor", "Lead Guest Editor", "Area Chair"):
            assert f"{role_type}: " in text and f"{role_type} - " not in text
    else:
        assert all(t not in text for t in ("Organizing Committee", "Paper Selection Committee", "IPMI 2027", "Associate Editor", "Area Chair", "Polo Tecnologico"))
        assert [p.paragraph_format.space_after for p in doc.paragraphs] == [Pt(3.5), Pt(0)], "The short signature needs a 3.5pt gap before the motto"

    link_nodes = [node for p in doc.paragraphs for node in p._p.iter(qn("w:hyperlink"))]
    assert len(link_nodes) == n_links, f"Expected {n_links} hyperlinks"
    relations = {rel.target_ref for rel in doc.part.rels.values() if rel.reltype.endswith("/hyperlink")}
    assert "mailto:raffaele.mineo@phd.unict.it" in relations
    assert "https://ingraffaelemineo.github.io/" in relations
    displayed_links = ["".join(t.text or "" for t in node.iter(qn("w:t"))) for node in link_nodes]
    assert "https://ingraffaelemineo.github.io/" in displayed_links, "Display the full HTTPS URL with trailing slash"
    assert "ingraffaelemineo.github.io" not in displayed_links, "Bare domain should not appear in signature hyperlinks"

    seen_name = False
    for p in doc.paragraphs:
        assert p.paragraph_format.space_before in (None, Pt(0))
        assert p.paragraph_format.space_after in (None, Pt(0), Pt(3.5))
        assert abs(p.paragraph_format.line_spacing - 1.0) < 0.001
        for r in p._p.iter(qn("w:r")):
            part_text = "".join(t.text or "" for t in r.iter(qn("w:t")))
            rPr = r.find(qn("w:rPr"))
            if rPr is None or not part_text:
                continue
            f = rPr.find(qn("w:rFonts"))
            assert f is not None and f.get(qn("w:ascii")) == "Segoe UI", part_text
            b = rPr.find(qn("w:b"))
            is_bold = b is not None and b.get(qn("w:val")) != "0"
            assert is_bold == (part_text == "Raffaele Mineo"), f"Unexpected bold: {part_text}"
            if is_bold:
                seen_name = True
            italic = rPr.find(qn("w:i"))
            is_italic = italic is not None and italic.get(qn("w:val")) != "0"
            if part_text == motto:
                assert is_italic, "Motto must preserve its understated italic treatment"
            elif "🙃" in part_text:
                assert not is_italic, "Emoji must remain upright"
            else:
                assert not is_italic, f"Only motto may be italic: {part_text}"
            size = rPr.find(qn("w:sz"))
            assert size is not None and int(size.get(qn("w:val"))) == (20 if is_bold else 18)
            color = rPr.find(qn("w:color"))
            assert color is not None
            expected_color = "404040" if is_bold else ("365F7D" if r.getparent().tag == qn("w:hyperlink") else "606060")
            assert color.get(qn("w:val")) == expected_color, (part_text, expected_color)
            underline = rPr.find(qn("w:u"))
            if underline is not None:
                assert underline.get(qn("w:val")) == "none", part_text
    assert seen_name
    print(f"PASS {path.name}: {n_paragraphs} paragraphs, {n_soft_breaks} soft breaks, {n_links} links")


if __name__ == "__main__":
    check(FULL, 4, 9, 2, full=True)
    check(SHORT, 2, 3, 2, full=False)
